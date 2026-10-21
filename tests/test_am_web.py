import io
import docx
import pytest
from fastapi.testclient import TestClient

import am_web, server, sdf, sessions

PEOPLE = [{"id": "tri", "name": "Le Minh Tri", "department": "DEV", "role": "employee", "expected_hours": 4},
          {"id": "vi", "name": "Vi Nguyen", "department": None, "role": "employee", "expected_hours": 7},
          {"id": "thang", "name": "Thang Dam Quyet", "department": "SEO", "role": "employee", "expected_hours": 7}]
TASKS = [{"id": "t1", "name": "Fix header", "status": "In Progress", "project_id": "p1", "pic_id": "tri",
          "pic_estimation": 10, "pic_start_date": "2026-10-19", "pic_due_date": "2026-10-23"},
         {"id": "t2", "name": "Stuck job", "status": "Stuck", "project_id": "p1", "pic_id": "vi",
          "pic_estimation": 4, "pic_start_date": "2026-10-26", "pic_due_date": "2026-10-26"}]


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("WUT_DEMO", "1")                              # passes the sign-in gate
    monkeypatch.setattr(sessions, "DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setattr(am_web, "OUT_DIR", str(tmp_path / "out"))
    monkeypatch.setattr(sdf, "fetch_tasks", lambda s, e: TASKS)
    monkeypatch.setattr(sdf, "fetch_people", lambda: PEOPLE)
    monkeypatch.setattr(sdf, "fetch_projects", lambda: {"p1": "Kenno"})
    return TestClient(server.app), tmp_path


def gen(c, week="2026-10-19"):
    return c.post("/api/am-report/generate", json={"week": week})


def test_generate_returns_blanks_and_pictures(client):
    c, tmp = client
    r = gen(c)
    assert r.status_code == 200
    body = r.json()
    assert body["week_start"] == "2026-10-19" and body["data"]["date_text"] == "23/10/26"
    assert "Scope Creep & Upsell Tracking" in body["blank"]
    assert "Kenno: 1 in progress (Fix header)" in body["data"]["this_week"]
    assert body["data"]["blockers"] == ["Stuck job (Kenno): Stuck"]
    assert len(body["pictures"]) == 6                      # 3 matched people x (this week, next week)
    assert body["warnings"] and any("Nguyen Nguyen Thi My" in w for w in body["warnings"])   # roster name with no match
    for p in body["pictures"]:
        assert c.get(p["url"]).status_code == 200
    text = r.text
    assert str(tmp) not in text and "@" not in text        # no server paths, no emails


def test_idle_people_are_flagged_for_next_week(client):
    c, _ = client
    body = gen(c).json()
    assert "Thang Dam Quyet" in body["idle"]               # no hours next week
    assert "Le Minh Tri" in body["idle"]                   # his task ends this Friday


def test_not_configured_returns_503_and_writes_nothing(client, monkeypatch):
    c, tmp = client
    def boom(*a):
        raise sdf.SdfNotConfigured("SD Friday isn't connected. Add SDF_READONLY_DB_URL to .env.")
    monkeypatch.setattr(sdf, "fetch_people", boom)
    r = gen(c)
    assert r.status_code == 503 and "SDF_READONLY_DB_URL" in r.json()["detail"]
    assert not (tmp / "out").exists()


def test_database_errors_do_not_leak_details(client, monkeypatch):
    c, _ = client
    def boom(*a):
        raise RuntimeError("connection to server at host db.secret.example password=hunter2 failed")
    monkeypatch.setattr(sdf, "fetch_tasks", boom)
    r = gen(c)
    assert r.status_code == 502 and "secret" not in r.text and "hunter2" not in r.text


def test_picture_route_rejects_traversal_and_unknown_names(client):
    c, _ = client
    gen(c)
    good = "/api/am-report/pictures/2026-10-19/"
    assert c.get(good + "nope.png").status_code == 404
    for bad in ("..%2Fsettings.json", "%2e%2e%2f%2e%2e%2fetc%2fpasswd", "x.png%00"):
        assert c.get(good + bad).status_code in (400, 404)
    assert c.get("/api/am-report/pictures/..%2F..%2Fdata/x.png").status_code in (400, 404)
    assert c.get("/api/am-report/pictures/not-a-date/dev-1-this.png").status_code in (400, 404)


def test_export_returns_a_docx_with_the_edited_text(client):
    c, _ = client
    data = gen(c).json()
    data["data"]["this_week"] = ["Edited line"]
    r = c.post("/api/am-report/export", json={"week_start": data["week_start"], "data": data["data"]})
    assert r.status_code == 200 and "wordprocessingml" in r.headers["content-type"]
    assert "Account%20Manager%20Report" in r.headers["content-disposition"]
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
    assert "Edited line" in text and "Kenno: 1 in progress" not in text


def test_export_rejects_unknown_picture_names(client):
    c, _ = client
    data = gen(c).json()["data"]
    data["dev_pictures"] = ["../../etc/passwd.png"]
    assert c.post("/api/am-report/export", json={"week_start": "2026-10-19", "data": data}).status_code == 400


def test_non_monday_week_is_normalised(client):
    c, _ = client
    assert gen(c, "2026-10-21").json()["week_start"] == "2026-10-19"


def test_bad_week_is_a_400(client):
    c, _ = client
    assert gen(c, "yesterday").status_code == 400


def test_empty_week_still_builds(client, monkeypatch):
    c, _ = client
    monkeypatch.setattr(sdf, "fetch_tasks", lambda s, e: [])
    body = gen(c).json()
    assert body["data"]["last_week"] == [] and body["data"]["blockers"] == []
    assert c.post("/api/am-report/export", json={"week_start": body["week_start"], "data": body["data"]}).status_code == 200


def test_settings_round_trip(client):
    c, _ = client
    s = c.get("/api/am-report/settings").json()
    s["submitter"] = "Mimi"
    assert c.put("/api/am-report/settings", json=s).json()["submitter"] == "Mimi"
    assert c.put("/api/am-report/settings", json={**s, "submitter": ""}).status_code == 400


def test_template_endpoint_lists_every_prompt(client):
    c, _ = client
    t = c.get("/api/am-report/template").json()
    keys = [i["key"] for s in t["sections"] for sub in s["subsections"] for i in sub["items"]]
    assert len(keys) == 12 and "scope_creep" in keys and t["title"].startswith("Account Manager")


def test_export_carries_typed_answers(client):
    c, _ = client
    r = c.post("/api/am-report/export", json={"week_start": "2026-10-19", "data": {"answers": {"scope_creep": "Typed by hand"}}})
    assert r.status_code == 200
    assert "Typed by hand" in "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)


def test_overlong_or_unknown_answers_are_rejected(client):
    c, _ = client
    big = {"answers": {"scope_creep": "x" * 6000}}
    assert c.post("/api/am-report/export", json={"week_start": "2026-10-19", "data": big}).status_code == 400
    odd = {"answers": {"not_a_prompt": "x"}}
    assert c.post("/api/am-report/export", json={"week_start": "2026-10-19", "data": odd}).status_code == 400


def test_picture_urls_change_on_every_generate_and_are_not_cached(client):
    c, _ = client
    first = gen(c).json()["pictures"][0]["url"]
    second = gen(c).json()["pictures"][0]["url"]
    assert "?v=" in first and first != second               # the browser must refetch after "Generate again"
    assert c.get(first).headers["cache-control"] == "no-store"
