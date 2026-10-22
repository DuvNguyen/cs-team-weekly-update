# pipeline.py — run.py split in two so a UI can sit between "generate" and "publish".
# generate: per client, fetch GSC data, draw the chart, write the summary (nothing is created in Google).
# publish:  create the Google Doc from the (possibly edited) previews.
# Set WUT_DEMO=1 to run the whole flow on fake data, without Google or Gemini.
import datetime, io, json, os, re, threading, time, uuid
from dotenv import load_dotenv

load_dotenv()   # .env next to this file; real environment variables win

import auth, auth_web, chart, dates, gsc, report, trends, updates, voice, voice_store
from googleapiclient.discovery import build
from google import genai

DEMO = os.environ.get("WUT_DEMO") == "1"
CONFIG_PATH = os.environ.get("WUT_CONFIG", "config.json")
OUT_DIR = "out"

_voice_lock = threading.Lock()    # report.build_prompt reads the voice from module globals, so swap + build is one step
_api_lock = threading.Lock()      # googleapiclient/httplib2 clients are not thread-safe
_svc: dict = {}
_svc_lock = threading.Lock()
JOBS: dict = {}


def load_config() -> dict:
    with open(CONFIG_PATH) as f:
        config = json.load(f)
    config["gemini_api_key"] = os.environ.get("GEMINI_API_KEY", config.get("gemini_api_key", ""))
    return config


def public_config() -> dict:
    c = load_config()
    return {
        "clients": c["clients"],
        "seo_sheet_tab": c.get("seo_sheet_tab", ""),
        "include_industry_trends": c.get("include_industry_trends", False),
        "demo": DEMO,
    }


def save_clients(clients: list) -> None:
    with open(CONFIG_PATH) as f:
        config = json.load(f)
    config["clients"] = clients
    with open(CONFIG_PATH, "w") as f:
        json.dump(config, f, indent=2, ensure_ascii=False)


def services(config: dict, email: str | None = None):
    """Google + Gemini clients. With an email: that signed-in user's tokens. Without: the old desktop flow."""
    with _svc_lock:
        key = email or ""
        if key not in _svc:
            if email:
                creds = auth_web.credentials_for(email)
                _svc[key] = {
                    "webmasters": build("searchconsole", "v1", credentials=creds),
                    "sheets": build("sheets", "v4", credentials=creds),
                    "docs": build("docs", "v1", credentials=creds),
                    "drive": build("drive", "v3", credentials=creds),
                    "gemini": genai.Client(api_key=config["gemini_api_key"]),
                }
            else:
                _svc[key] = auth.get_services(config)       # opens a browser for Google sign-in
        return _svc[key]


def forget(email: str) -> None:
    with _svc_lock:
        _svc.pop(email, None)


auth_web.on_forget(forget)


def list_tabs(email: str | None = None) -> list:
    if DEMO:
        return ["14/9-25/9-New", "28/9-9/10", "12/10-23/10"]
    config = load_config()
    svc = services(config, email)
    meta = svc["sheets"].spreadsheets().get(
        spreadsheetId=config["seo_sheet_id"], fields="sheets.properties.title").execute()
    return [s["properties"]["title"] for s in meta.get("sheets", [])]


_latest: dict = {}      # email -> (fetched_at, iso date)


def latest_data_date(email: str | None = None) -> str:
    """Newest day Search Console has data for, which is where its own "7 days" ends.
    The lag changes (2-3 days), so ask the API instead of assuming. Tries clients until one is readable."""
    if DEMO:
        return default_period()["end"]
    hit = _latest.get(email)
    if hit and time.time() - hit[0] < 1800:
        return hit[1]
    config = load_config()
    svc = services(config, email)
    today = datetime.date.today()
    last_err = None
    for c in config["clients"]:
        try:
            with _api_lock:
                rows = svc["webmasters"].searchanalytics().query(siteUrl=c["property"], body={
                    "startDate": (today - datetime.timedelta(days=10)).isoformat(), "endDate": today.isoformat(),
                    "dimensions": ["date"], "type": "web"}).execute().get("rows", [])
            if rows:
                day = max(r["keys"][0] for r in rows)
                _latest[email] = (time.time(), day)
                return day
        except Exception as e:       # no access to this property: try the next client
            last_err = e
    raise RuntimeError(f"Could not read Search Console: {last_err}")


def default_period() -> dict:
    start, end = dates.report_period(datetime.date.today())
    return {"start": start.isoformat(), "end": end.isoformat()}


_DATE_HEADER = re.compile(r"^\s*\d{1,2}/\d{1,2}\s*[-–]\s*\d{1,2}/\d{1,2}\s*:?\s*$")                  # "12/10 - 16/10"
_NO_WORK = re.compile(r"^\s*(no\s+tasks?|none|n/?a|nothing|tbd|-+|(haven'?t|not)\s+signed(\s+yet)?)\s*\.?\s*$", re.I)


def clean_notes(text: str) -> str:
    """The sheet cell holds one section per week under a date header. Drop sections with no real work
    ("No task", blank lines, "Haven't signed yet") so Gemini is never handed a placeholder as if it were work."""
    sections, head, body = [], None, []

    def flush():
        lines = [ln.rstrip() for ln in body if ln.strip() and not _NO_WORK.match(ln)]
        if lines:
            sections.append("\n".join(([head] if head else []) + lines))

    for line in (text or "").splitlines():
        if _DATE_HEADER.match(line):
            flush()
            head, body = line.strip(), []
        else:
            body.append(line)
    flush()
    return "\n\n".join(sections)


def _periods(start: datetime.date, end: datetime.date):
    days = (end - start).days + 1
    pend = start - datetime.timedelta(days=1)
    return pend - datetime.timedelta(days=days - 1), pend


def _demo_series(seed: int, start: datetime.date, end: datetime.date) -> dict:
    n = (end - start).days + 1
    rows = [{"keys": [(start + datetime.timedelta(days=i)).isoformat()],
             "clicks": 20 + (seed * 7 + i * 3) % 25,
             "impressions": 900 + (seed * 53 + i * 41) % 400,
             "position": 8 + (seed % 5) + (i % 3) * 0.4} for i in range(n)]
    return gsc.parse(rows)


def _prompt(style: dict, name: str, category: str, series: dict, prev: dict, this_week: str, next_week: str,
            start: datetime.date, end: datetime.date) -> str:
    """report.build_prompt, but written in the chosen style (stays unchanged, reads voice.VOICE/EXAMPLES/STRUCTURE)."""
    with _voice_lock:
        voice.VOICE, voice.EXAMPLES = voice_store.render(style)
        voice.STRUCTURE = voice_store.render_structure(style, start, end)
        return report.build_prompt(name, category, series, prev, this_week, next_week)


def _generate(gemini, prompt: str) -> str:
    return (gemini.models.generate_content(model=report.MODEL, contents=prompt).text or "").strip()


def new_job(tab: str, client_names: list, start: str, end: str, email: str | None = None, style_id: str | None = None) -> dict:
    config = load_config()
    style = voice_store.get_style(style_id)
    by_name = {c["name"]: c for c in config["clients"]}
    clients = []
    for name in client_names:
        c = by_name[name]
        clients.append({"name": name, "property": c["property"], "category": c.get("category", ""),
                        "status": "pending", "error": "", "include": True,
                        "totals": None, "prev_totals": None, "body": "", "has_chart": False,
                        "style_id": style["id"], "style_name": style["name"]})
    job = {"id": uuid.uuid4().hex[:10], "status": "queued", "message": "Queued",
           "tab": tab, "start": start, "end": end, "clients": clients,
           "doc_url": "", "error": "", "_rows": [], "_trends": {}, "_email": email}
    JOBS[job["id"]] = job
    threading.Thread(target=_run, args=(job,), daemon=True).start()
    return job


def view(job: dict) -> dict:
    return {k: v for k, v in job.items() if not k.startswith("_")}


def _run(job: dict) -> None:
    try:
        config = load_config()
        svc = None
        if not DEMO:
            job["status"], job["message"] = "signing_in", "Connecting to Google…"
            svc = services(config, job["_email"])
        job["status"], job["message"] = "running", "Reading SEO sheet…"
        _load_sheet(job, config, svc)
        _load_trends(job, config, svc)
        os.makedirs(OUT_DIR, exist_ok=True)
        for i, c in enumerate(job["clients"], 1):
            job["message"] = f"Client {i}/{len(job['clients'])}: {c['name']}"
            generate_client(job, c, config, svc)
        job["status"], job["message"] = "ready", "Review the previews, then publish."
    except Exception as e:
        job["status"], job["error"], job["message"] = "error", str(e), "Failed"


def _load_sheet(job, config, svc):
    sheet_id, tab = config.get("seo_sheet_id", ""), job["tab"].strip()
    job["_rows"] = []
    if DEMO or not (sheet_id and tab):
        return
    try:
        with _api_lock:
            job["_rows"] = updates.fetch_rows(svc["sheets"], sheet_id, f"'{tab}'!A1:C300")
    except Exception as e:
        job["message"] = f"Tab '{tab}' not found — continuing with charts only. ({str(e).splitlines()[0][:80]})"


def _load_trends(job, config, svc):
    if DEMO or not config.get("include_industry_trends", False):
        return
    manual = config.get("category_trends", {})
    for cat in sorted({c["category"] for c in job["clients"] if c["category"]}):
        if manual.get(cat):
            job["_trends"][cat] = (manual[cat], [])
        else:
            with _api_lock:
                job["_trends"][cat] = trends.for_category(svc["gemini"], cat, report.MODEL)


def generate_client(job: dict, c: dict, config: dict, svc) -> None:
    """Fill one client's preview. Safe to call again for Retry / Regenerate."""
    c["status"], c["error"] = "running", ""
    try:
        start, end = datetime.date.fromisoformat(job["start"]), datetime.date.fromisoformat(job["end"])
        pstart, pend = _periods(start, end)
        cfg = next((x for x in config["clients"] if x["name"] == c["name"]), {})
        style = voice_store.get_style(c.get("style_id"))
        c["style_id"], c["style_name"] = style["id"], style["name"]
        key = cfg.get("sheet_key", c["name"])
        this_week, next_week = updates.match(job["_rows"], key)
        this_week, next_week = cfg.get("this_week", this_week), cfg.get("next_week", next_week)
        this_week, next_week = clean_notes(this_week), clean_notes(next_week)
        if DEMO:
            seed = sum(map(ord, c["name"]))
            series, prev = _demo_series(seed, start, end), _demo_series(seed + 1, pstart, pend)
            summary = (f"Demo summary for {c['name']} in the style “{style['name']}”: clicks and impressions moved against the previous "
                       "period while our average position held steady. Next, our focus is on the content queue.")
            this_week = this_week or "Published two blog posts."
            next_week = next_week or "Refresh the top landing pages."
        else:
            with _api_lock:
                series = gsc.fetch(svc["webmasters"], c["property"], start, end)
                prev = gsc.fetch(svc["webmasters"], c["property"], pstart, pend)
            summary = None
        os.makedirs(OUT_DIR, exist_ok=True)
        chart.render(series, f"{c['name']} — {c['property']}", chart_path(job, c["name"]))
        if summary is None:
            prompt = _prompt(style, c["name"], c["category"], series, prev, this_week, next_week, start, end)
            with _api_lock:
                summary = _generate(svc["gemini"], prompt)
        trend, sources = job["_trends"].get(c["category"], ("", []))
        ctx = job.setdefault("_ctx", {})[c["name"]] = dict(series=series, prev=prev, this_week=this_week,
                                                            next_week=next_week, trend=trend, sources=sources)
        c["body"] = _email(config, cfg, c, summary, ctx)
        c["totals"], c["prev_totals"] = series["totals"], prev["totals"]
        c["has_chart"], c["status"] = True, "ok"
    except Exception as e:
        c["status"], c["error"] = "failed", str(e).splitlines()[0][:1000]
        c["include"] = False


def _email(config: dict, cfg: dict, c: dict, summary: str, ctx: dict) -> str:
    """The client email: greeting, the summary + SEO notes (compose_body), sign-off."""
    structured = voice_store.has_structure(voice_store.get_style(c.get("style_id")))     # the sections already cover the work
    body = report.compose_body(summary, ctx["this_week"], ctx["next_week"], c["category"], ctx["trend"], ctx["sources"], notes=not structured)
    hi = cfg.get("contact") or f"{c['name']} team"
    sender = config.get("sender", "Mimi Pham")
    return f"Hi {hi},\n\n{body}\n\nBest regards,\n{sender}\nSaigon Digital"


def regenerate_client(job: dict, name: str, instruction: str, style_id: str | None = None) -> None:
    """Rewrite one client's email from the stored numbers, steered by the CS person's prompt."""
    c = next(x for x in job["clients"] if x["name"] == name)
    ctx = job.get("_ctx", {}).get(name)
    if not ctx:
        raise ValueError("Nothing to regenerate yet for this client.")
    config = load_config()
    cfg = next((x for x in config["clients"] if x["name"] == name), {})
    instruction = instruction.strip()
    if style_id:
        c["style_id"] = style_id
    style = voice_store.get_style(c.get("style_id"))
    c["style_id"], c["style_name"] = style["id"], style["name"]
    if DEMO:
        summary = f"Demo summary for {name} in the style “{style['name']}”, rewritten with your instruction: {instruction or 'none'}."
    else:
        prompt = _prompt(style, name, c["category"], ctx["series"], ctx["prev"], ctx["this_week"], ctx["next_week"],
                         datetime.date.fromisoformat(job["start"]), datetime.date.fromisoformat(job["end"]))
        if instruction:
            prompt += f"\n\nExtra instruction from the account manager (follow it, keep the rules above): {instruction}"
        svc = services(config, job["_email"])
        with _api_lock:
            summary = _generate(svc["gemini"], prompt)
    c["body"] = _email(config, cfg, c, summary, ctx)


def chart_path(job: dict, name: str) -> str:
    safe = "".join(ch if ch.isalnum() else "_" for ch in name)
    return os.path.join(OUT_DIR, f"{job['id']}_{safe}.png")


def retry_client(job: dict, name: str) -> None:
    c = next(x for x in job["clients"] if x["name"] == name)
    config = load_config()
    svc = None if DEMO else services(config, job["_email"])
    generate_client(job, c, config, svc)


def publish(job: dict) -> str:
    chosen = [c for c in job["clients"] if c["include"] and c["status"] == "ok"]
    if not chosen:
        raise ValueError("No clients selected to publish.")
    job["status"], job["message"] = "publishing", "Creating the Google Doc…"
    try:
        if DEMO:
            job["doc_url"] = "https://docs.google.com/document/d/demo/edit"
        else:
            config = load_config()
            svc = services(config, job["_email"])
            title = f"Bi-weekly SEO Report — {job['start']} to {job['end']}"
            with _api_lock:
                doc_id = report.create_doc(svc["docs"], svc["drive"], title, config["drive_folder_id"])
                for i, c in enumerate(chosen, 1):
                    job["message"] = f"Adding {c['name']} ({i}/{len(chosen)})"
                    report.append_client(svc["docs"], svc["drive"], doc_id, c["name"],
                                         chart_path(job, c["name"]), c["body"],
                                         config["drive_folder_id"], c["category"])
            job["doc_url"] = report.doc_url(doc_id)
        job["status"], job["message"] = "done", "Report created."
        return job["doc_url"]
    except Exception as e:
        job["status"], job["error"], job["message"] = "ready", str(e), "Publish failed — you can retry."
        raise


# ---------- Word export (local download, no Google Docs needed) ----------

def _add_hyperlink(paragraph, url: str, text: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r_id = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "1155CC")
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single")
    props.append(color); props.append(underline); run.append(props)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    run.append(t); link.append(run)
    paragraph._p.append(link)


def _add_text_with_links(paragraph, line: str) -> None:
    pos = 0
    for m in report.URL_RE.finditer(line):
        url = m.group(0).rstrip(".,);]")
        end = m.start() + len(url)
        if m.start() > pos:
            paragraph.add_run(line[pos:m.start()])
        _add_hyperlink(paragraph, url, url)
        pos = end
    if pos < len(line):
        paragraph.add_run(line[pos:])


def export_docx(job: dict) -> tuple:
    """Build the report as a Word file from the previews on screen. Returns (bytes, filename)."""
    from docx import Document
    from docx.shared import Inches
    chosen = [c for c in job["clients"] if c["include"] and c["status"] == "ok"]
    if not chosen:
        raise ValueError("No clients selected to export.")
    doc = Document()
    doc.add_heading(f"Bi-weekly SEO Report — {job['start']} to {job['end']}", level=0)
    for i, c in enumerate(chosen):
        if i:
            doc.add_page_break()
        doc.add_heading(c["name"], level=1)
        if c["category"]:
            doc.add_paragraph(c["category"])
        doc.add_picture(chart_path(job, c["name"]), width=Inches(6.3))
        for line in c["body"].split("\n"):
            _add_text_with_links(doc.add_paragraph(), line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), f"SEO-Report_{job['start']}_to_{job['end']}.docx"


def analyze_structure(email_text: str) -> dict:
    """Propose the email's sections for review (nothing is saved). The email stays the source; every excerpt is checked against it."""
    import structure_draft
    if DEMO:
        return structure_draft.paragraphs(email_text)
    return structure_draft.analyze(genai.Client(api_key=load_config()["gemini_api_key"]), report.MODEL, email_text)


def analyze_style(email_text: str) -> dict:
    """Draft a writing style from a pasted email (nothing is saved)."""
    if DEMO:
        return {"id": "demo", "name": "Demo style", "persona": "Write as an account manager sending a client update.",
                "style": [{"id": "d1", "text": "Keep it short and friendly.", "enabled": True}],
                "avoid": [{"id": "d2", "text": "Jargon.", "enabled": True}],
                "examples": [{"id": "d3", "label": "the email this style came from", "text": email_text.strip(), "enabled": True}],
                "structure": [{"id": "d4", "title": "Performance overview", "instruction": "Clicks, impressions and position with meaning.", "enabled": True},
                              {"id": "d5", "title": "Next steps ({next_period})", "instruction": "A bulleted list of the plan.", "enabled": True}]}
    return voice_store.analyze(genai.Client(api_key=load_config()["gemini_api_key"]), report.MODEL, email_text)
