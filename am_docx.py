# am_docx.py — builds the Account Manager End-of-Week Report as a styled .docx.
# Wording, order and bullets come from am_template.json (copied from the sample .docx, never retyped);
# only the look is new. Anything we can't fill is an empty yellow-highlighted line, never invented text.
import json
import os
import re
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TEMPLATE = os.path.join(os.path.dirname(__file__), "am_template.json")
NAVY, BLUE, GREY, WASH = RGBColor(0x06, 0x15, 0x28), RGBColor(0x31, 0x65, 0xF5), RGBColor(0x5F, 0x6B, 0x7A), "EBF0FE"
FONT = "Calibri"
PICTURE_WIDTH = Cm(15)


def _clean(text) -> str:
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(text or ""))   # control characters are illegal in .docx XML


def _shade(el_pr, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


def _bottom_border(paragraph, colour: str = "3165F5", size: int = 8) -> None:
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", str(size)), ("space", "2"), ("color", colour)):
        bottom.set(qn(f"w:{k}"), v)
    pbdr.append(bottom)
    paragraph._p.get_or_add_pPr().append(pbdr)


def _run(paragraph, text, *, bold=False, italic=False, size=None, colour=None):
    r = paragraph.add_run(_clean(text))
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if colour is not None:
        r.font.color.rgb = colour
    return r


def _setup(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    for name, size, colour, bold in (("Normal", 10.5, None, False), ("Title", 22, NAVY, True),
                                     ("Heading 1", 15, NAVY, True), ("Heading 2", 12.5, BLUE, True),
                                     ("List Bullet", 10.5, None, False), ("List Bullet 2", 10.5, None, False)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = FONT, Pt(size), bold
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        if colour is not None:
            st.font.color.rgb = colour
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)


def _blank(doc, style="Normal") -> None:
    """An empty answer line: highlighted so the CS team can't miss it."""
    p = doc.add_paragraph(style=style)
    r = p.add_run(" " * 70)
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    _run(p, text, bold=True)
    if level == 1:
        _bottom_border(p)


def _header_block(doc, template, data) -> None:
    rows = [(template["header"]["submitted_by"], data.get("submitter", "")),
            (template["header"]["date"], f'{data.get("date_text", "")} {template["header"]["due_note"]}'.strip()),
            (template["header"]["reviewer"], data.get("reviewer", ""))]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row, (label, value) in zip(table.rows, rows):
        row.cells[0].width, row.cells[1].width = Cm(5), Cm(12)
        for cell in row.cells:
            _shade(cell._tc.get_or_add_tcPr(), WASH)
        _run(row.cells[0].paragraphs[0], label, bold=True, colour=NAVY)
        _run(row.cells[1].paragraphs[0], value)


def _pictures(doc, paths) -> None:
    for path in paths:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.add_run().add_picture(path, width=PICTURE_WIDTH)


def _bullets(doc, lines) -> None:
    for line in lines:
        _run(doc.add_paragraph(style="List Bullet 2"), line)


def _signatures(doc, template) -> None:
    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=len(template["signatures"]))
    for cell, sig in zip(table.rows[0].cells, template["signatures"]):
        p = cell.paragraphs[0]
        _run(p, sig["name"], bold=True, colour=NAVY)
        _run(cell.add_paragraph(), sig["role"], colour=GREY)
        for _ in range(3):
            cell.add_paragraph()
        line = cell.add_paragraph()
        _bottom_border(line, "9AA7BB", 4)
        _run(cell.add_paragraph(), "Date: ", colour=GREY)


def build(out_path: str, data: dict) -> list[str]:
    """Write the report. Returns the labels of the prompts left blank."""
    template = json.load(open(TEMPLATE, encoding="utf-8"))
    doc = Document()
    _setup(doc)
    title = doc.add_paragraph(style="Title")
    _run(title, template["title"], bold=True)
    _bottom_border(title, size=12)
    _header_block(doc, template, data)
    blanks: list[str] = []

    for section in template["sections"]:
        _heading(doc, section["heading"], 1)
        for sub in section["subsections"]:
            if sub["heading"]:
                _heading(doc, sub["heading"], 2)
            _run(doc.add_paragraph(), sub["objective"], italic=True, size=10, colour=GREY)
            for item in sub["items"]:
                prompt = doc.add_paragraph(style="List Bullet")
                prompt.paragraph_format.space_before = Pt(8)
                prompt.paragraph_format.keep_with_next = True
                _run(prompt, item["label"], bold=True, colour=NAVY)
                _run(prompt, f': {item["text"]}')
                answer = item["answer"]
                if answer == "client_lists":
                    lists = [data.get("last_week") or [], data.get("this_week") or []]
                    for heading, lines in zip(item["lists"], lists):
                        p = doc.add_paragraph()
                        _run(p, heading, bold=True)
                        if lines:
                            _bullets(doc, lines)
                        else:
                            _blank(doc, "List Bullet 2")
                    filled = any(lists)
                elif answer in ("dev_pictures", "seo_pictures"):
                    filled = bool(data.get(answer))
                    _pictures(doc, data.get(answer) or []) if filled else _blank(doc)
                elif answer == "blockers":
                    filled = bool(data.get("blockers"))
                    _bullets(doc, data["blockers"]) if filled else _blank(doc)
                else:
                    filled = False
                    _blank(doc)
                if not filled:
                    blanks.append(item["label"])
    _signatures(doc, template)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    return blanks
